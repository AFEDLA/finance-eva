"""
skills/quotation/quotation_generator.py
Generate Quotation PDF & Excel — layout persis template Holomoc.

Logo & tanda tangan di-ekstrak dari Quotation_Template_AI.xlsx:
  image1.png → Logo HOLOMOC (464×60px) — dipakai di header PDF & Excel
  image2.png → Tanda tangan Chandra Kirana (442×186px) — dipakai di TTD PDF & Excel
"""

import os
import zipfile
from pathlib import Path
from datetime import datetime

BASE_DIR   = Path(os.getenv("STORAGE_DIR", str(Path.home() / "Documents" / "holomoc-file")))
EXPORT_DIR = BASE_DIR / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

_SKILLS_DIR = Path(os.path.dirname(os.path.abspath(__file__)))

COMPANY_NAME    = "PT Holomoc Indonesia"
COMPANY_ADDRESS = "Jl. Duren Sawit Indah No.K1 no 10,"
HOLOMOC_BLUE    = "1B6EC2"
BLACK_HEX       = "000000"
WHITE_HEX       = "FFFFFF"


# ─── LOGO & SIGNATURE RESOLVER ────────────────────────────────────────────────

def _resolve_assets():
    """
    Return (logo_path, signature_path) — keduanya Path atau None.
    Cari di skills/quotation/ dulu, fallback ekstrak dari template xlsx.
    """
    logo_path = _SKILLS_DIR / "qt_image1.png"
    sig_path  = _SKILLS_DIR / "qt_image2.png"

    # Kalau belum ada, coba ekstrak dari template
    if not logo_path.exists() or not sig_path.exists():
        candidates = [
            _SKILLS_DIR / "Quotation_Template_AI.xlsx",
            _SKILLS_DIR.parent / "Quotation_Template_AI.xlsx",
            BASE_DIR / "Quotation_Template_AI.xlsx",
        ]
        for tpl in candidates:
            if tpl.exists():
                try:
                    with zipfile.ZipFile(str(tpl)) as z:
                        if not logo_path.exists():
                            logo_path.write_bytes(z.read("xl/media/image1.png"))
                        if not sig_path.exists():
                            sig_path.write_bytes(z.read("xl/media/image2.png"))
                    break
                except Exception:
                    pass

    return (
        logo_path  if logo_path.exists() else None,
        sig_path   if sig_path.exists()  else None,
    )


# ─── HELPERS ──────────────────────────────────────────────────────────────────

def _fmt_date(tanggal: str) -> str:
    bulan_map = {
        "Januari":"January","Februari":"February","Maret":"March",
        "April":"April","Mei":"May","Juni":"June","Juli":"July",
        "Agustus":"August","September":"September","Oktober":"October",
        "November":"November","Desember":"December"
    }
    parts = tanggal.strip().split()
    if len(parts) == 3:
        day, mon, year = parts
        return f"{bulan_map.get(mon, mon)} {day}, {year}"
    return tanggal


# ─── EXCEL ────────────────────────────────────────────────────────────────────

def generate_qt_excel(qt: dict) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XLImage

    logo_path, sig_path = _resolve_assets()

    items   = qt.get("items", [])
    total   = sum(int(x.get("jumlah", 0)) for x in items)
    tanggal = _fmt_date(qt.get("tanggal", datetime.now().strftime("%B %d, %Y")))
    qt_id   = qt.get("id", "QT")

    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws.sheet_view.showGridLines = False

    # Lebar kolom persis template
    ws.column_dimensions["A"].width = 6.3
    ws.column_dimensions["B"].width = 68.0
    ws.column_dimensions["C"].width = 6.7
    ws.column_dimensions["D"].width = 6.0
    ws.column_dimensions["E"].width = 18.3
    ws.column_dimensions["F"].width = 32.4

    def F(bold=False, size=11, color="000000", name="Calibri"):
        return Font(name=name, bold=bold, size=size, color=color)

    def A(h="left", v="center", wrap=False):
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

    thin = Side(style="thin",  color="000000")
    none = Side(style=None)

    def border(top=None, bottom=None, left=None, right=None):
        return Border(
            top=top or none, bottom=bottom or none,
            left=left or none, right=right or none
        )

    black_fill = PatternFill("solid", fgColor=BLACK_HEX)

    # ═══ ROW 1 kosong ══════════════════════════════════════════════════════════
    ws.row_dimensions[1].height = 8

    # ═══ ROW 2 — Logo HOLOMOC + Date ═══════════════════════════════════════════
    ws.row_dimensions[2].height = 30
    if logo_path:
        try:
            from PIL import Image as PILImage
            with PILImage.open(str(logo_path)) as pil:
                w_px, h_px = pil.size
            # Target tinggi ~1.1cm (row height 30pt ≈ 40px)
            # Excel row height dalam point (1pt ≈ 1.33px)
            target_h_px = 38
            scale = target_h_px / h_px
            target_w_px = int(w_px * scale)
            img = XLImage(str(logo_path))
            img.width  = target_w_px
            img.height = target_h_px
            img.anchor = "A2"
            ws.add_image(img)
        except Exception:
            c = ws.cell(row=2, column=1, value="HOLOMOC")
            c.font = Font(name="Calibri", bold=True, size=24, color=HOLOMOC_BLUE)
            c.alignment = A("left")
    else:
        c = ws.cell(row=2, column=1, value="HOLOMOC")
        c.font = Font(name="Calibri", bold=True, size=24, color=HOLOMOC_BLUE)
        c.alignment = A("left")

    ws.cell(row=2, column=5, value="Date  :").font      = F(size=11)
    ws.cell(row=2, column=5).alignment                   = A("right")
    ws.cell(row=2, column=6, value=tanggal).font         = F(size=11)
    ws.cell(row=2, column=6).alignment                   = A("right")

    # ═══ ROW 3 — PT Holomoc Indonesia + No ═════════════════════════════════════
    ws.cell(row=3, column=1, value=COMPANY_NAME).font      = F(size=11)
    ws.cell(row=3, column=1).alignment                      = A("left")
    ws.cell(row=3, column=5, value="No  :").font            = F(size=11)
    ws.cell(row=3, column=5).alignment                      = A("right")
    ws.cell(row=3, column=6, value=qt.get("nomor","")).font = F(size=11)
    ws.cell(row=3, column=6).alignment                      = A("right")

    # ═══ ROW 4 — Alamat + UP ═══════════════════════════════════════════════════
    ws.row_dimensions[4].height = 16.9
    ws.cell(row=4, column=1, value=COMPANY_ADDRESS).font = F(size=11)
    ws.cell(row=4, column=1).alignment                    = A("left")

    klien_up   = qt.get("klien_up", "")
    klien_nama = qt.get("klien_nama", "")
    up_val     = klien_up if klien_up else klien_nama

    ws.cell(row=4, column=5, value="UP:").font  = F(size=11)
    ws.cell(row=4, column=5).alignment           = A("right")
    ws.cell(row=4, column=6, value=up_val).font  = F(size=11)
    ws.cell(row=4, column=6).alignment           = A("right")

    # ═══ ROW 5 — Nama klien bold kanan ═════════════════════════════════════════
    ws.row_dimensions[5].height = 17.45
    ws.cell(row=5, column=1, value=qt.get("klien_alamat","")).font      = F(size=11)
    ws.cell(row=5, column=1).alignment                                   = A("left")
    ws.cell(row=5, column=6, value=klien_nama).font                      = F(bold=True, size=11)
    ws.cell(row=5, column=6).alignment                                   = A("right")

    # ═══ ROW 6, 7 kosong ═══════════════════════════════════════════════════════
    ws.row_dimensions[6].height = 16.9

    # ═══ ROW 8 — Re (merge A:F, center bold size 16) ═══════════════════════════
    ws.row_dimensions[8].height = 21.0
    ws.merge_cells("A8:F8")
    c8 = ws.cell(row=8, column=1, value=qt.get("re",""))
    c8.font      = Font(name="Calibri", bold=True, size=16)
    c8.alignment = A("center")

    # ═══ ROW 9 — Header tabel HITAM ════════════════════════════════════════════
    ws.row_dimensions[9].height = 24.0

    def hdr(col, val):
        c = ws.cell(row=9, column=col, value=val)
        c.font      = Font(name="Calibri", bold=True, color=WHITE_HEX, size=12)
        c.fill      = black_fill
        c.alignment = A("center")
        return c

    hdr(1, "No"); hdr(2, "Description"); hdr(3, "Qty")
    hdr(4, ""); hdr(5, "Price/IDR"); hdr(6, "Total/IDR")
    ws.merge_cells("C9:D9")

    ws.cell(row=9, column=1).border = border(top=thin, bottom=thin, left=thin)
    for col in range(2, 6):
        ws.cell(row=9, column=col).border = border(top=thin, bottom=thin)
    ws.cell(row=9, column=6).border = border(top=thin, bottom=thin, right=thin)

    # ═══ ROW 10+ — Item rows ═══════════════════════════════════════════════════
    row = 10
    for i, item in enumerate(items):
        qty    = item.get("qty", 1)
        sat    = item.get("satuan", "Ls")
        harga  = int(item.get("harga_per_unit", item.get("harga", 0)))
        jumlah = int(item.get("jumlah", harga * qty))
        desc   = item.get("deskripsi", item.get("nama", ""))
        is_first = (i == 0)

        ws.row_dimensions[row].height = 15

        for col, val, aln, bold, fmt in [
            (1, i+1,    A("center"), True,  None),
            (2, desc,   A("left", wrap=True), True, None),
            (3, qty,    A("right"), False, None),
            (4, sat,    A("left"),  False, None),
            (5, harga,  A("right"), False, '#,##0'),
            (6, jumlah, A("right"), False, '#,##0'),
        ]:
            c = ws.cell(row=row, column=col, value=val)
            c.font      = F(bold=bold, size=11)
            c.alignment = aln
            c.border    = border(top=thin, bottom=thin,
                                 left=thin if col == 1 else thin,
                                 right=thin if col == 6 else thin)
            if fmt and isinstance(val, int): c.number_format = fmt
        row += 1

    # Baris kosong dalam tabel
    ws.row_dimensions[row].height = 15
    for col in range(1, 7):
        ws.cell(row=row, column=col).border = border(top=thin, bottom=thin, left=thin, right=thin)
    row += 1

    # ═══ Total row — HITAM ═════════════════════════════════════════════════════
    ws.row_dimensions[row].height = 18
    ws.merge_cells(f"A{row}:E{row}")
    # Isi merge cell via kolom pertama (A), kolom lain sudah ter-merge
    c_total_lbl = ws.cell(row=row, column=1, value="Total")
    c_total_lbl.font      = Font(name="Calibri", size=11, color=WHITE_HEX)
    c_total_lbl.fill      = black_fill
    c_total_lbl.alignment = A("left")
    c_total_lbl.border    = border(top=thin, bottom=none, left=thin)
    # Fill hitam untuk kolom F
    c = ws.cell(row=row, column=6, value=total)
    c.font          = Font(name="Calibri", size=11, color=WHITE_HEX)
    c.fill          = black_fill
    c.alignment     = A("right")
    c.number_format = '#,##0'
    c.border        = border(top=thin, bottom=none, right=thin)
    row += 1

    # ═══ Note ══════════════════════════════════════════════════════════════════
    if qt.get("note"):
        ws.row_dimensions[row].height = 15
        ws.cell(row=row, column=1, value="Note :").font      = F(size=11)
        ws.cell(row=row, column=1).alignment                  = A("left")
        ws.cell(row=row, column=2, value=qt["note"]).font     = F(size=11)
        ws.cell(row=row, column=2).alignment                  = A("left", wrap=True)
        row += 1

    # ═══ TTD — Regards (F), gambar signature (F), nama (F), HP (F) ═════════════
    # Semua di kolom F saja (col 6), tidak merge agar aman dengan image anchor
    ttd_nama = qt.get("ttd_nama", "Chandra Kirana")
    ttd_hp   = qt.get("ttd_hp",   "081285232001")

    row += 1  # spasi

    # Regards
    ws.row_dimensions[row].height = 16
    c = ws.cell(row=row, column=6, value="Regards")
    c.font      = F(bold=True, size=11)
    c.alignment = A("center")
    row += 1

    # Gambar signature — anchor di kolom F baris ini
    # Target tinggi 2cm ≈ 56px pada 72dpi screen
    SIG_H_PX = 56
    if sig_path:
        try:
            from PIL import Image as PILImage
            with PILImage.open(str(sig_path)) as pil:
                w_px, h_px = pil.size
            scale       = SIG_H_PX / h_px
            target_w_px = int(w_px * scale)

            sig_img        = XLImage(str(sig_path))
            sig_img.width  = target_w_px
            sig_img.height = SIG_H_PX
            sig_img.anchor = f"F{row}"
            ws.add_image(sig_img)
        except Exception:
            pass

    # Baris untuk ruang gambar (4 baris × 14pt ≈ 75px)
    for _ in range(4):
        ws.row_dimensions[row].height = 14
        row += 1

    # Nama TTD
    ws.row_dimensions[row].height = 16
    c = ws.cell(row=row, column=6, value=ttd_nama)
    c.font      = F(bold=True, size=11)
    c.alignment = A("center")
    row += 1

    # HP TTD
    ws.row_dimensions[row].height = 14
    c = ws.cell(row=row, column=6, value=ttd_hp)
    c.font      = F(bold=True, size=11)
    c.alignment = A("center")

    fname = f"QT_{qt_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path  = EXPORT_DIR / fname
    wb.save(str(path))
    return str(path)


# ─── PDF ──────────────────────────────────────────────────────────────────────

def _fmt_rp(n: int) -> str:
    return f"{n:,}".replace(",", ".")

def generate_qt_pdf(qt: dict) -> str:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                    Paragraph, Spacer, Image as RLImage)
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_CENTER

    logo_path, sig_path = _resolve_assets()

    items   = qt.get("items", [])
    total   = sum(int(x.get("jumlah", 0)) for x in items)
    tanggal = _fmt_date(qt.get("tanggal", datetime.now().strftime("%B %d, %Y")))
    qt_id   = qt.get("id", "QT")

    fname = f"QT_{qt_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    path  = EXPORT_DIR / fname

    doc = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=1.8*cm, rightMargin=1.8*cm,
        topMargin=1.5*cm,  bottomMargin=1.5*cm
    )

    BLUE_C  = colors.HexColor("#1B6EC2")
    BLACK_C = colors.HexColor("#000000")
    GREY_C  = colors.HexColor("#555555")

    s_co    = ParagraphStyle("co",  fontSize=9,  fontName="Helvetica")
    s_lbl   = ParagraphStyle("lbl", fontSize=9,  fontName="Helvetica", alignment=TA_RIGHT)
    s_sep   = ParagraphStyle("sep", fontSize=9,  fontName="Helvetica", alignment=TA_CENTER)
    s_val   = ParagraphStyle("val", fontSize=9,  fontName="Helvetica", alignment=TA_RIGHT)
    s_kl    = ParagraphStyle("kl",  fontSize=9,  fontName="Helvetica", spaceAfter=1)
    s_re    = ParagraphStyle("re",  fontSize=13, fontName="Helvetica-Bold", alignment=TA_CENTER, spaceAfter=4)
    s_cell  = ParagraphStyle("tc",  fontSize=9)
    s_bold  = ParagraphStyle("tb",  fontSize=9,  fontName="Helvetica-Bold")
    s_rg    = ParagraphStyle("rg",  fontSize=9,  fontName="Helvetica-Bold", alignment=TA_CENTER)
    s_ttd   = ParagraphStyle("ttd", fontSize=9,  fontName="Helvetica-Bold", alignment=TA_CENTER)
    s_hp    = ParagraphStyle("hp",  fontSize=9,  alignment=TA_CENTER)

    klien_up   = qt.get("klien_up", "")
    klien_nama = qt.get("klien_nama", "")
    up_val     = klien_up if klien_up else klien_nama

    story = []

    # ── Header: Logo kiri + Date/No/UP kanan ──────────────────────────────────
    if logo_path:
        try:
            from PIL import Image as PILImage
            with PILImage.open(str(logo_path)) as pil:
                w_px, h_px = pil.size
            # Logo Holomoc: target tinggi 1.2cm, width proporsional dari ratio pixel
            logo_target_h = 1.2 * cm
            logo_w = logo_target_h * w_px / h_px
            left_content = [
                RLImage(str(logo_path), width=logo_w, height=logo_target_h),
                Spacer(1, 0.1*cm),
                Paragraph(COMPANY_NAME,    s_co),
                Paragraph(COMPANY_ADDRESS, s_co),
            ]
        except Exception:
            left_content = _text_logo_col(s_co)
    else:
        left_content = [
            Paragraph("<b><font color='#1B6EC2' size=20>HOLOMOC</font></b>",
                      ParagraphStyle("lg", fontSize=20, fontName="Helvetica-Bold")),
            Spacer(1, 0.1*cm),
            Paragraph(COMPANY_NAME,    s_co),
            Paragraph(COMPANY_ADDRESS, s_co),
        ]

    info_rows = [("Date", tanggal), ("No", qt.get("nomor","")), ("UP", up_val)]
    right_data = [[Paragraph(l, s_lbl), Paragraph(":", s_sep), Paragraph(v, s_val)]
                  for l, v in info_rows]
    right_inner = Table(right_data, colWidths=[1.5*cm, 0.4*cm, 6.5*cm])
    right_inner.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("TOPPADDING",(0,0),(-1,-1),1),("BOTTOMPADDING",(0,0),(-1,-1),2),
        ("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),0),
    ]))
    hdr_tbl = Table([[left_content, right_inner]], colWidths=[9*cm, 8.7*cm])
    hdr_tbl.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),0),
        ("TOPPADDING",(0,0),(-1,-1),0),("BOTTOMPADDING",(0,0),(-1,-1),0),
    ]))
    story.append(hdr_tbl)
    story.append(Spacer(1, 0.4*cm))

    if klien_nama: story.append(Paragraph(klien_nama, s_kl))
    if qt.get("klien_alamat"): story.append(Paragraph(qt["klien_alamat"], s_kl))
    story.append(Spacer(1, 0.4*cm))
    story.append(Paragraph(qt.get("re",""), s_re))
    story.append(Spacer(1, 0.1*cm))

    # ── Tabel item ──────────────────────────────────────────────────────────────
    col_w = [0.8*cm, 8.5*cm, 1.3*cm, 1.3*cm, 3.0*cm, 3.0*cm]
    tbl_data = [["No", "Description", "Qty", "Sat.", "Price/IDR", "Total/IDR"]]
    for i, item in enumerate(items):
        qty    = item.get("qty", 1)
        sat    = item.get("satuan", "Ls")
        harga  = int(item.get("harga_per_unit", item.get("harga", 0)))
        jumlah = int(item.get("jumlah", harga * qty))
        desc   = item.get("deskripsi", item.get("nama", ""))
        is_first = (i == 0)
        tbl_data.append([
            Paragraph(str(i+1), s_cell),
            Paragraph(desc, s_bold),
            Paragraph(str(qty), s_cell),
            Paragraph(sat, s_cell),
            Paragraph(_fmt_rp(harga) if harga else "", s_cell),
            Paragraph(_fmt_rp(jumlah), s_cell),
        ])
    tbl_data.append(["", Paragraph("Total", s_bold), "", "", "",
                      Paragraph(_fmt_rp(total), s_bold)])

    t = Table(tbl_data, colWidths=col_w, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0), BLACK_C),
        ("TEXTCOLOR",(0,0),(-1,0), colors.white),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,0),9),
        ("ALIGN",(0,0),(-1,0),"CENTER"),
        ("FONTSIZE",(0,1),(-1,-1),9),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("ALIGN",(0,1),(0,-1),"CENTER"),
        ("ALIGN",(1,1),(1,-1),"LEFT"),
        ("ALIGN",(2,1),(2,-1),"RIGHT"),
        ("ALIGN",(3,1),(3,-1),"LEFT"),
        ("ALIGN",(4,1),(4,-1),"RIGHT"),
        ("ALIGN",(5,1),(5,-1),"RIGHT"),
        ("BACKGROUND",(0,-1),(-1,-1), BLACK_C),
        ("TEXTCOLOR",(0,-1),(-1,-1), colors.white),
        ("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold"),
        ("ALIGN",(5,-1),(5,-1),"RIGHT"),
        ("BOX",(0,0),(-1,-1), 0.5, BLACK_C),
        ("INNERGRID",(0,0),(-1,-1), 0.3, colors.HexColor("#AAAAAA")),
        ("TOPPADDING",(0,0),(-1,-1),4),
        ("BOTTOMPADDING",(0,0),(-1,-1),4),
        ("LEFTPADDING",(1,0),(1,-1),4),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.3*cm))

    if qt.get("note"):
        story.append(Paragraph(f"Note :  {qt['note']}", s_cell))
        story.append(Spacer(1, 0.2*cm))
    story.append(Spacer(1, 0.6*cm))

    # ── TTD: tabel 2 kolom — kiri kosong, kanan Regards + sig + nama + HP ────────
    ttd_nama = qt.get("ttd_nama", "Chandra Kirana")
    ttd_hp   = qt.get("ttd_hp",   "081285232001")

    # Ukuran signature — proporsional, max width 4cm
    sig_w = 4.0 * cm
    sig_h_val = 1.8 * cm  # default kalau PIL gagal
    if sig_path:
        try:
            from PIL import Image as PILImage
            with PILImage.open(str(sig_path)) as pil:
                wpx, hpx = pil.size
            sig_h_val = sig_w * hpx / wpx
        except Exception:
            pass

    # Buat signature flowable
    if sig_path:
        try:
            sig_flowable = RLImage(str(sig_path), width=sig_w, height=sig_h_val)
        except Exception:
            sig_flowable = Spacer(1, sig_h_val)
    else:
        sig_flowable = Spacer(1, sig_h_val)

    # Kolom kanan: semua elemen TTD dalam satu list
    right_col = [
        Paragraph("Regards", s_rg),
        Spacer(1, 0.15*cm),
        sig_flowable,
        Spacer(1, 0.1*cm),
        Paragraph(f"<b>{ttd_nama}</b>", s_ttd),
        Paragraph(ttd_hp, s_hp),
    ]

    # Tabel TTD: 2 kolom, kiri kosong, kanan konten
    ttd_tbl = Table(
        [[Spacer(1, 0.1*cm), right_col]],
        colWidths=[9.5*cm, 8.2*cm]
    )
    ttd_tbl.setStyle(TableStyle([
        ("VALIGN",        (0,0), (-1,-1), "TOP"),
        ("ALIGN",         (1,0), (1,0),   "CENTER"),
        ("LEFTPADDING",   (0,0), (-1,-1), 0),
        ("RIGHTPADDING",  (0,0), (-1,-1), 0),
        ("TOPPADDING",    (0,0), (-1,-1), 0),
        ("BOTTOMPADDING", (0,0), (-1,-1), 0),
    ]))
    story.append(ttd_tbl)
    doc.build(story)
    return str(path)
