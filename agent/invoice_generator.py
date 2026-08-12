"""
agent/invoice_generator.py — EVA Invoice Document Generator
============================================================
Generate PDF dan Excel Invoice dari data storage.
Template referensi: Invoice_Template_AI.xlsx
"""

import os
from pathlib import Path
from datetime import datetime

# ─── PATH SETUP ───────────────────────────────────────────────────────────────
_THIS_DIR   = Path(__file__).parent
_ASSETS_DIR = _THIS_DIR / "assets"
_LOGO_PATH  = _ASSETS_DIR / "inv_logo.png"

def _get_output_dir() -> Path:
    out = Path(os.getenv("STORAGE_DIR", str(Path.home() / "Documents" / "holomoc-file"))) / "exports"
    out.mkdir(parents=True, exist_ok=True)
    return out

def _fmt_rp(n) -> str:
    try:
        return f"Rp {int(n):,.0f}".replace(",", ".")
    except:
        return "Rp 0"

def _total_items(items: list) -> int:
    return sum(int(x.get("total", x.get("jumlah", 0))) for x in items)

def _terbilang(n: int) -> str:
    if n == 0:
        return "Nol Rupiah"
    satuan = ["", "Satu", "Dua", "Tiga", "Empat", "Lima",
              "Enam", "Tujuh", "Delapan", "Sembilan", "Sepuluh",
              "Sebelas", "Dua Belas", "Tiga Belas", "Empat Belas",
              "Lima Belas", "Enam Belas", "Tujuh Belas", "Delapan Belas", "Sembilan Belas"]
    def _tiga(x):
        if x == 0: return ""
        if x < 20: return satuan[x]
        if x < 100:
            return satuan[x // 10] + " Puluh" + ("" if x % 10 == 0 else " " + satuan[x % 10])
        ratus = "Seratus" if x // 100 == 1 else satuan[x // 100] + " Ratus"
        sisa = x % 100
        return ratus + ("" if sisa == 0 else " " + _tiga(sisa))
    hasil = ""
    if n >= 1_000_000_000:
        hasil += _tiga(n // 1_000_000_000) + " Miliar "
        n %= 1_000_000_000
    if n >= 1_000_000:
        hasil += _tiga(n // 1_000_000) + " Juta "
        n %= 1_000_000
    if n >= 1_000:
        rb = n // 1_000
        hasil += ("Seribu " if rb == 1 else _tiga(rb) + " Ribu ")
        n %= 1_000
    if n > 0:
        hasil += _tiga(n)
    return hasil.strip() + " Rupiah"


# ═══════════════════════════════════════════════════════════════════════════════
# PDF GENERATOR
# ═══════════════════════════════════════════════════════════════════════════════

def generate_inv_pdf(inv: dict) -> str:
    """Generate PDF Invoice. Return path ke file yang dibuat."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm, mm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    )
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_CENTER

    inv_id   = inv.get("id", "INV-0000")
    out_path = _get_output_dir() / f"{inv_id}.pdf"

    # ── Page setup ──────────────────────────────────────────────────────────
    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        leftMargin=1.8*cm, rightMargin=1.8*cm,
        topMargin=1.5*cm,  bottomMargin=2*cm,
    )
    W = A4[0] - 3.6*cm   # usable width ≈ 493pt

    # ── Styles ──────────────────────────────────────────────────────────────
    def S(name, **kw):
        base = {"fontName": "Helvetica", "fontSize": 10, "leading": 14,
                "textColor": colors.black}
        base.update(kw)
        return ParagraphStyle(name, **base)

    sNormal  = S("normal")
    sBold    = S("bold",   fontName="Helvetica-Bold")
    sRight   = S("right",  alignment=TA_RIGHT)
    sBoldR   = S("boldr",  fontName="Helvetica-Bold", alignment=TA_RIGHT)
    sCenter  = S("ctr",    alignment=TA_CENTER)
    sBoldC   = S("boldC",  fontName="Helvetica-Bold", alignment=TA_CENTER)
    sSmallB  = S("smB",    fontSize=9, leading=12, fontName="Helvetica-Bold")
    sSmall   = S("sm",     fontSize=9, leading=12)
    sTitle   = S("title",  fontName="Helvetica-Bold", fontSize=22,
                           alignment=TA_CENTER, leading=28)
    sCompany = S("co",     fontSize=11, leading=15, fontName="Helvetica-Bold")
    sAddr    = S("addr",   fontSize=9,  leading=13)
    # Header label style — right-aligned, no extra leading
    sHdrLbl  = S("hLbl",  fontName="Helvetica-Bold", fontSize=10,
                          alignment=TA_RIGHT, leading=14)
    sHdrVal  = S("hVal",  fontName="Helvetica-Bold", fontSize=10,
                          alignment=TA_LEFT,  leading=14)

    story = []

    # ── HEADER: logo (kiri) + nomor/tanggal (kanan) ─────────────────────────
    # Struktur: [logo_cell | right_info]
    # right_info: 2 baris (Invoice: nomor, Date: tanggal)
    # Kolom nomor dan tanggal cukup lebar, tanpa nested table berlebih

    logo_cell = Spacer(1, 1)
    if _LOGO_PATH.exists():
        from reportlab.platypus import Image as RLImage
        # Logo image 476x84px, aspect ≈ 5.67
        # Tinggi disesuaikan dengan area header (≈1.4cm)
        logo_w = 6.0 * cm
        logo_h = logo_w / 5.667          # proporsi asli ≈ 1.06cm
        logo_cell = RLImage(str(_LOGO_PATH), width=logo_w, height=logo_h)

    nomor   = inv.get("nomor", "")
    tanggal = inv.get("tanggal", "")

    # Right-side info sebagai 2-row table dengan 3 kolom: label | titik dua | value
    # Lebar kolom safe — semua padding di-zero-kan
    col_lbl = 1.8*cm    # "Invoice" / "Date"
    col_sep = 0.5*cm    # ":"
    col_val = W*0.55 - col_lbl - col_sep  # sisa

    right_tbl = Table(
        [
            [Paragraph("Invoice", sHdrLbl), Paragraph(":", sHdrVal), Paragraph(nomor, sHdrVal)],
            [Paragraph("Date",    sHdrLbl), Paragraph(":", sHdrVal), Paragraph(tanggal, sHdrVal)],
        ],
        colWidths=[col_lbl, col_sep, col_val],
        rowHeights=[16, 16],
        style=[
            ("VALIGN",        (0,0), (-1,-1), "MIDDLE"),
            ("LEFTPADDING",   (0,0), (-1,-1), 0),
            ("RIGHTPADDING",  (0,0), (-1,-1), 0),
            ("TOPPADDING",    (0,0), (-1,-1), 0),
            ("BOTTOMPADDING", (0,0), (-1,-1), 2),
        ]
    )

    header_tbl = Table(
        [[logo_cell, right_tbl]],
        colWidths=[W*0.45, W*0.55],
        style=[
            ("VALIGN",       (0,0), (-1,-1), "TOP"),
            ("LEFTPADDING",  (0,0), (-1,-1), 0),
            ("RIGHTPADDING", (0,0), (-1,-1), 0),
            ("TOPPADDING",   (0,0), (-1,-1), 0),
            ("BOTTOMPADDING",(0,0), (-1,-1), 0),
        ]
    )
    story.append(header_tbl)
    story.append(Spacer(1, 3*mm))

    # ── COMPANY INFO ──────────────────────────────────────────────────────────
    story.append(Paragraph("PT Holomoc Indonesia", sCompany))
    story.append(Paragraph("Jl. Duren Sawit Indah No.K1 no 10,", sAddr))
    story.append(Paragraph("Kec. Duren Sawit, Jakarta Timur, 13440.", sAddr))
    story.append(Spacer(1, 5*mm))

    # ── JUDUL "INVOICE" + double underline ────────────────────────────────────
    story.append(Paragraph("INVOICE", sTitle))
    # Double underline: 2 HRFlowable dengan jarak kecil (seperti template Excel double border)
    story.append(HRFlowable(width="100%", thickness=1.0, color=colors.black,
                             spaceBefore=2, spaceAfter=1))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.black,
                             spaceBefore=0, spaceAfter=4))
    story.append(Spacer(1, 4*mm))

    # ── CLIENT SECTION ────────────────────────────────────────────────────────
    klien_nama   = inv.get("klien_nama", "")
    klien_alamat = inv.get("klien_alamat", "")
    project_name = inv.get("project_name", "")

    alamat_lines = [l.strip()
                    for l in klien_alamat.replace("\\n", "\n").split("\n")
                    if l.strip()]

    client_rows = [[Paragraph("To:", sBold), Paragraph(klien_nama, sBold)]]
    for al in alamat_lines:
        client_rows.append(["", Paragraph(al, sNormal)])
    if project_name:
        client_rows.append(["", ""])
        client_rows.append(["", Paragraph("Project Name:", sNormal)])
        client_rows.append(["", Paragraph(project_name, sNormal)])

    client_tbl = Table(client_rows, colWidths=[1.2*cm, W - 1.2*cm])
    client_tbl.setStyle(TableStyle([
        ("VALIGN",        (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING",   (0,0), (-1,-1), 0),
        ("RIGHTPADDING",  (0,0), (-1,-1), 0),
        ("TOPPADDING",    (0,0), (-1,-1), 1),
        ("BOTTOMPADDING", (0,0), (-1,-1), 1),
    ]))
    story.append(client_tbl)
    story.append(Spacer(1, 5*mm))

    # ── ITEM TABLE ────────────────────────────────────────────────────────────
    # Lebar kolom: No | Item | Satuan(Rp) | Qty | Price(Rp)
    col_no    = 1.0*cm
    col_item  = W * 0.40
    col_sat   = W * 0.18
    col_qty   = W * 0.13
    col_price = W - col_no - col_item - col_sat - col_qty

    items = inv.get("items", [])
    grand_total = _total_items(items)

    # Build ALL rows: header row + item rows + empty closer + total row
    n_items = len(items)
    all_rows = []
    styles   = []

    # ── Row 0: header ────────────────────────────────────────────────────────
    all_rows.append([
        Paragraph("No",          sBoldC),
        Paragraph("Item",        sBoldC),
        Paragraph("Satuan (Rp)", sBoldC),
        Paragraph("Qty",         sBoldC),
        Paragraph("Price (Rp)",  sBoldC),
    ])
    styles += [
        ("BACKGROUND",    (0,0), (-1,0), colors.white),
        ("BOX",           (0,0), (-1,0), 0.5, colors.black),
        ("INNERGRID",     (0,0), (-1,0), 0.5, colors.black),
        ("FONTNAME",      (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE",      (0,0), (-1,0), 10),
        ("ALIGN",         (0,0), (-1,0), "CENTER"),
        ("VALIGN",        (0,0), (-1,0), "MIDDLE"),
        ("TOPPADDING",    (0,0), (-1,0), 5),
        ("BOTTOMPADDING", (0,0), (-1,0), 5),
        ("LEFTPADDING",   (0,0), (-1,0), 3),
        ("RIGHTPADDING",  (0,0), (-1,0), 3),
    ]

    # ── Row 1..N: item rows ──────────────────────────────────────────────────
    for i, it in enumerate(items, 1):
        row_idx = i   # 0=header, i=item row
        nama     = it.get("nama", it.get("deskripsi", ""))
        qty      = it.get("qty", 1)
        satuan   = it.get("satuan", "Ls")
        sat_rp   = it.get("satuan_rp", 0)
        total_it = it.get("total", it.get("jumlah", 0))

        all_rows.append([
            Paragraph(str(i),                                   sCenter),
            Paragraph(nama,                                      sNormal),
            Paragraph(_fmt_rp(sat_rp)   if sat_rp   else "",    sRight),
            Paragraph(f"{qty} {satuan}",                         sCenter),
            Paragraph(_fmt_rp(total_it) if total_it else "",     sRight),
        ])
        styles += [
            # Full box around entire row
            ("BOX",           (0, row_idx), (-1, row_idx), 0.5, colors.black),
            # Inner vertical lines
            ("INNERGRID",     (0, row_idx), (-1, row_idx), 0.5, colors.black),
            ("VALIGN",        (0, row_idx), (-1, row_idx), "MIDDLE"),
            ("TOPPADDING",    (0, row_idx), (-1, row_idx), 4),
            ("BOTTOMPADDING", (0, row_idx), (-1, row_idx), 4),
            ("LEFTPADDING",   (0, row_idx), (-1, row_idx), 3),
            ("RIGHTPADDING",  (0, row_idx), (-1, row_idx), 3),
        ]

    # ── Row N+1: empty spacer row (bottom border of table body) ─────────────
    spacer_idx = n_items + 1
    all_rows.append(["", "", "", "", ""])
    styles += [
        ("TOPPADDING",    (0, spacer_idx), (-1, spacer_idx), 1),
        ("BOTTOMPADDING", (0, spacer_idx), (-1, spacer_idx), 1),
        # Left/right borders continuity
        ("LINEBEFORE",  (0, spacer_idx), (0, spacer_idx),  0.5, colors.black),
        ("LINEAFTER",   (-1, spacer_idx), (-1, spacer_idx), 0.5, colors.black),
        ("LINEBELOW",   (0, spacer_idx), (-1, spacer_idx),  0.5, colors.black),
    ]

    # ── Row N+2: Total row ──────────────────────────────────────────────────
    total_idx = n_items + 2
    all_rows.append([
        "", "",
        Paragraph("Total", sBoldR),
        "",
        Paragraph(_fmt_rp(grand_total), sBoldR),
    ])
    styles += [
        ("SPAN",          (2, total_idx), (3, total_idx)),
        ("BOX",           (4, total_idx), (4, total_idx), 0.5, colors.black),
        ("LINEABOVE",     (2, total_idx), (-1, total_idx), 0.5, colors.black),
        ("VALIGN",        (0, total_idx), (-1, total_idx), "MIDDLE"),
        ("TOPPADDING",    (0, total_idx), (-1, total_idx), 5),
        ("BOTTOMPADDING", (0, total_idx), (-1, total_idx), 5),
        ("LEFTPADDING",   (0, total_idx), (-1, total_idx), 3),
        ("RIGHTPADDING",  (0, total_idx), (-1, total_idx), 3),
    ]

    tbl = Table(
        all_rows,
        colWidths=[col_no, col_item, col_sat, col_qty, col_price],
    )
    tbl.setStyle(TableStyle(styles))
    story.append(tbl)
    story.append(Spacer(1, 5*mm))

    # ── TERBILANG ──────────────────────────────────────────────────────────────
    story.append(Paragraph("Terbilang:", sSmallB))
    story.append(Paragraph(f"#{_terbilang(grand_total)},--", sSmallB))
    story.append(Spacer(1, 7*mm))

    # ── FOOTER: BANK INFO (kiri) + TTD (kanan) ───────────────────────────────
    # Gunakan satu row dengan dua kolom — masing-masing kolom berisi
    # nested Table 1-kolom sehingga bank info dan TTD tidak saling
    # mempengaruhi tinggi baris satu sama lain.
    from reportlab.platypus import KeepInFrame

    bank_nama = inv.get("bank_nama", "")
    bank_rek  = inv.get("bank_rekening", "")
    bank_cab  = inv.get("bank_cabang", "")
    bank_an   = inv.get("bank_atas_nama", "")
    ttd_nama  = inv.get("ttd_nama", "")
    ttd_jab   = inv.get("ttd_jabatan", "Finance")

    W_bank = W * 0.55
    W_ttd  = W * 0.45

    # ── Kiri: bank info sebagai nested single-col table ──────────────────────
    bank_rows = []
    if bank_nama or bank_rek:
        bank_rows.append([Paragraph("Pembayaran Melalui Rekening:", sSmallB)])
        line2 = " ".join(filter(None, [
            bank_nama,
            f"Rekening: {bank_rek}" if bank_rek else "",
        ]))
        if line2:
            bank_rows.append([Paragraph(line2, sSmallB)])
        if bank_cab:
            bank_rows.append([Paragraph(f"Cabang {bank_cab}", sSmallB)])
        if bank_an:
            bank_rows.append([Paragraph(f"a/n {bank_an}", sSmallB)])

    if bank_rows:
        bank_inner = Table(bank_rows, colWidths=[W_bank])
        bank_inner.setStyle(TableStyle([
            ("VALIGN",        (0,0), (-1,-1), "TOP"),
            ("LEFTPADDING",   (0,0), (-1,-1), 0),
            ("RIGHTPADDING",  (0,0), (-1,-1), 4),
            ("TOPPADDING",    (0,0), (-1,-1), 2),
            ("BOTTOMPADDING", (0,0), (-1,-1), 2),
        ]))
        bank_cell = bank_inner
    else:
        bank_cell = Paragraph("", sSmall)

    # ── Kanan: TTD sebagai nested single-col table ───────────────────────────
    ttd_rows = [
        [Paragraph("Prepared by,", sCenter)],
        [Spacer(1, 18*mm)],
        [Paragraph(ttd_nama if ttd_nama else "________________", sCenter)],
        [Paragraph(ttd_jab, sCenter)],
    ]
    ttd_inner = Table(ttd_rows, colWidths=[W_ttd])
    ttd_inner.setStyle(TableStyle([
        ("ALIGN",         (0,0), (-1,-1), "CENTER"),
        ("VALIGN",        (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING",   (0,0), (-1,-1), 0),
        ("RIGHTPADDING",  (0,0), (-1,-1), 0),
        ("TOPPADDING",    (0,0), (-1,-1), 2),
        ("BOTTOMPADDING", (0,0), (-1,-1), 2),
    ]))

    # ── Gabung dalam satu outer row ──────────────────────────────────────────
    footer_tbl = Table(
        [[bank_cell, ttd_inner]],
        colWidths=[W_bank, W_ttd],
    )
    footer_tbl.setStyle(TableStyle([
        ("VALIGN",        (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING",   (0,0), (-1,-1), 0),
        ("RIGHTPADDING",  (0,0), (-1,-1), 0),
        ("TOPPADDING",    (0,0), (-1,-1), 0),
        ("BOTTOMPADDING", (0,0), (-1,-1), 0),
    ]))
    story.append(footer_tbl)

    doc.build(story)
    return str(out_path)


# ═══════════════════════════════════════════════════════════════════════════════
# EXCEL GENERATOR
# ═══════════════════════════════════════════════════════════════════════════════

def generate_inv_excel(inv: dict) -> str:
    """Generate Excel Invoice. Return path ke file yang dibuat."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage

    inv_id   = inv.get("id", "INV-0000")
    out_path = _get_output_dir() / f"{inv_id}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = inv_id

    # ── Column widths (sesuai template) ────────────────────────────────────
    ws.column_dimensions["A"].width = 3.25
    ws.column_dimensions["B"].width = 4.75
    ws.column_dimensions["C"].width = 80.625
    ws.column_dimensions["D"].width = 28.75
    ws.column_dimensions["E"].width = 4.875
    ws.column_dimensions["F"].width = 10.5
    ws.column_dimensions["G"].width = 29.125
    ws.column_dimensions["H"].width = 1.5

    # ── Row heights (sesuai template) ──────────────────────────────────────
    row_heights = {
        2: 35.1, 3: 18.75, 4: 18.75, 5: 18.75, 6: 17.25, 7: 32.25,
        8: 2.45, 9: 18.75, 10: 17.25, 11: 17.25, 12: 6.0,
        13: 17.25, 14: 17.25, 15: 12.75,
    }
    for r, h in row_heights.items():
        ws.row_dimensions[r].height = h

    # ── Style helpers ───────────────────────────────────────────────────────
    thin   = Side(style="thin")
    double = Side(style="double")

    def mk_font(bold=False, size=13, name="Calibri"):
        return Font(name=name, bold=bold, size=size)

    def mk_align(h="left", v="center", wrap=False):
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

    def mk_border(top=None, bottom=None, left=None, right=None):
        return Border(top=top, bottom=bottom, left=left, right=right)

    def set_cell(ws, coord, value, bold=False, size=13, halign="left",
                 border_obj=None, font_name="Calibri"):
        c = ws[coord]
        c.value     = value
        c.font      = mk_font(bold=bold, size=size, name=font_name)
        c.alignment = mk_align(h=halign, v="center")
        if border_obj:
            c.border = border_obj

    # ── ROW 2: Logo ────────────────────────────────────────────────────────
    ws.row_dimensions[2].height = 35.1
    if _LOGO_PATH.exists():
        try:
            xl_logo = XLImage(str(_LOGO_PATH))
            # Logo image: 476x84px, aspect ≈ 5.667
            # Row 2 height = 35.1pt → in pixels at 96dpi: 35.1*(96/72) ≈ 46.8px
            # Set height to match row exactly, width proportional
            logo_h_px = 46          # match row 2 height (35.1pt ≈ 46px at 96dpi)
            logo_w_px = round(logo_h_px * 5.667)  # ≈ 261px
            xl_logo.width  = logo_w_px
            xl_logo.height = logo_h_px
            xl_logo.anchor = "B2"
            ws.add_image(xl_logo)
        except Exception:
            pass

    # ── ROW 3: Company name + Invoice label + Nomor ─────────────────────────
    set_cell(ws, "B3", "PT Holomoc Indonesia", bold=False, size=14)
    set_cell(ws, "D3", "Invoice",              bold=True,  size=13, halign="right")
    set_cell(ws, "F3", ":",                    bold=True,  size=13)
    set_cell(ws, "G3", inv.get("nomor",""),    bold=True,  size=13, halign="right")

    # ── ROW 4: Alamat + Date ────────────────────────────────────────────────
    set_cell(ws, "B4", "Jl. Duren Sawit Indah No.K1 no 10,", bold=False, size=14)
    set_cell(ws, "D4", "Date",                bold=True, size=13, halign="right")
    set_cell(ws, "F4", ":",                   bold=True, size=13)
    set_cell(ws, "G4", inv.get("tanggal",""), bold=True, size=13, halign="right")

    # ── ROW 5: Alamat baris 2 ───────────────────────────────────────────────
    set_cell(ws, "B5", "Kec. Duren Sawit, Jakarta Timur, 13440.", bold=False, size=14)

    # ── ROW 7: "INVOICE" judul besar + double bottom border ─────────────────
    ws.merge_cells("B7:G7")
    c7             = ws["B7"]
    c7.value       = "INVOICE"
    c7.font        = mk_font(bold=True, size=24)
    c7.alignment   = mk_align(h="center", v="center")
    c7.border      = mk_border(bottom=double)

    # ── ROW 9-11: Client info ───────────────────────────────────────────────
    set_cell(ws, "B9", "To:", bold=False, size=13)
    set_cell(ws, "C9", inv.get("klien_nama",""), bold=True, size=13)

    klien_alamat = inv.get("klien_alamat", "")
    alamat_lines = [l.strip()
                    for l in klien_alamat.replace("\\n", "\n").split("\n")
                    if l.strip()]
    for i, al in enumerate(alamat_lines[:2]):
        set_cell(ws, f"C{10+i}", al, bold=False, size=13)

    # ── ROW 13-14: Project name ──────────────────────────────────────────────
    set_cell(ws, "C13", "Project Name:",           bold=False, size=13)
    set_cell(ws, "C14", inv.get("project_name",""), bold=False, size=13)

    # ── ROW 16: Table header ─────────────────────────────────────────────────
    ws.row_dimensions[16].height = 17.25
    ws.merge_cells("E16:F16")

    tbl_border = mk_border(top=thin, bottom=thin, left=thin, right=thin)
    set_cell(ws, "B16", "No",         bold=True, size=13, halign="center", border_obj=tbl_border)
    set_cell(ws, "C16", "Item",       bold=True, size=13, halign="center", border_obj=tbl_border)
    set_cell(ws, "D16", "Satuan (Rp)",bold=True, size=13, halign="center", border_obj=tbl_border)

    ws["E16"].value     = "Qty"
    ws["E16"].font      = mk_font(bold=True, size=13)
    ws["E16"].alignment = mk_align(h="center", v="center")
    ws["E16"].border    = tbl_border

    set_cell(ws, "G16", "Price (Rp)", bold=True, size=13, halign="center", border_obj=tbl_border)

    # ── ROW 17+: Items ──────────────────────────────────────────────────────
    items     = inv.get("items", [])
    row_start = 17
    item_h    = 18.95

    for i, it in enumerate(items):
        r    = row_start + i
        ws.row_dimensions[r].height = item_h

        nama     = it.get("nama", it.get("deskripsi", ""))
        qty      = it.get("qty", 1)
        satuan   = it.get("satuan", "Ls")
        sat_rp   = it.get("satuan_rp", 0)
        total_it = it.get("total", it.get("jumlah", 0))

        lr_border = mk_border(left=thin, right=thin)

        # B: No
        c = ws.cell(row=r, column=2)
        c.value     = i + 1
        c.font      = mk_font(size=13)
        c.alignment = mk_align(h="center", v="center")
        c.border    = lr_border

        # C: Item name
        c = ws.cell(row=r, column=3)
        c.value     = nama
        c.font      = mk_font(size=14)
        c.alignment = mk_align(h="left", v="center")

        # D: Satuan (Rp)
        c = ws.cell(row=r, column=4)
        c.value     = sat_rp if sat_rp else None
        c.font      = mk_font(size=13)
        c.alignment = mk_align(h="right", v="center")
        c.border    = lr_border
        if sat_rp:
            c.number_format = '#,##0'

        # E: Qty
        c = ws.cell(row=r, column=5)
        c.value     = qty
        c.font      = mk_font(size=13)
        c.alignment = mk_align(h="center", v="center")
        c.border    = lr_border

        # F: Satuan unit
        c = ws.cell(row=r, column=6)
        c.value     = satuan
        c.font      = mk_font(size=13)
        c.alignment = mk_align(h="center", v="center")
        c.border    = lr_border

        # G: Price total
        c = ws.cell(row=r, column=7)
        c.value     = total_it if total_it else None
        c.font      = mk_font(size=13)
        c.alignment = mk_align(h="right", v="center")
        c.border    = lr_border
        if total_it:
            c.number_format = '#,##0'

    # ── Bottom border row ────────────────────────────────────────────────────
    r_bot = row_start + len(items)
    ws.row_dimensions[r_bot].height = 6.0
    for col in range(2, 8):
        c = ws.cell(row=r_bot, column=col)
        c.border = mk_border(
            bottom=thin,
            left=thin  if col in (2, 5) else None,
            right=thin if col in (2, 5, 7) else None,
        )

    # ── Total row ────────────────────────────────────────────────────────────
    r_total = r_bot + 1
    ws.row_dimensions[r_total].height = 18.75
    ws.merge_cells(f"D{r_total}:F{r_total}")

    total_all = _total_items(items)
    c = ws[f"D{r_total}"]
    c.value     = "Total "
    c.font      = mk_font(bold=True, size=14)
    c.alignment = mk_align(h="right", v="center")
    c.border    = mk_border(top=thin, right=thin)

    c = ws[f"G{r_total}"]
    c.value         = total_all
    c.font          = mk_font(bold=True, size=14)
    c.alignment     = mk_align(h="right", v="center")
    c.border        = mk_border(top=thin, bottom=thin, left=thin, right=thin)
    c.number_format = '#,##0'

    # ── Terbilang ─────────────────────────────────────────────────────────────
    r_terb = r_total + 1
    ws.row_dimensions[r_terb].height = 17.25
    set_cell(ws, f"B{r_terb}", "Terbilang:", bold=True, size=13)

    r_terb2 = r_terb + 1
    ws.row_dimensions[r_terb2].height = 17.25
    set_cell(ws, f"B{r_terb2}", f"#{_terbilang(total_all)},--", bold=True, size=13)

    # ── TTD: Prepared by ──────────────────────────────────────────────────────
    ttd_nama    = inv.get("ttd_nama", "")
    ttd_jabatan = inv.get("ttd_jabatan", "Finance")

    r_ttd1 = r_terb2 + 1
    ws.row_dimensions[r_ttd1].height = 17.25
    set_cell(ws, f"G{r_ttd1}", "Prepared by,", bold=False, size=13, halign="center")

    for ri in range(r_ttd1 + 1, r_ttd1 + 5):
        ws.row_dimensions[ri].height = 17.25

    r_ttd_nama = r_ttd1 + 5
    ws.row_dimensions[r_ttd_nama].height = 17.25
    set_cell(ws, f"G{r_ttd_nama}", ttd_nama or "________________",
             bold=False, size=13, halign="center")

    r_ttd_jab = r_ttd_nama + 1
    ws.row_dimensions[r_ttd_jab].height = 17.25
    set_cell(ws, f"G{r_ttd_jab}", ttd_jabatan, bold=False, size=13, halign="center")

    # ── Bank info ─────────────────────────────────────────────────────────────
    bank_nama = inv.get("bank_nama", "")
    bank_rek  = inv.get("bank_rekening", "")
    bank_cab  = inv.get("bank_cabang", "")
    bank_an   = inv.get("bank_atas_nama", "")

    r_bank = r_ttd1 + 2
    if bank_nama or bank_rek:
        set_cell(ws, f"B{r_bank}",   "Pembayaran Melalui Rekening:", bold=True, size=13)
        line2 = " ".join(filter(None, [bank_nama, f"Rekening: {bank_rek}" if bank_rek else ""]))
        set_cell(ws, f"B{r_bank+1}", line2, bold=True, size=13)
        if bank_cab:
            set_cell(ws, f"B{r_bank+2}", f"Cabang {bank_cab}", bold=True, size=13)
        if bank_an:
            set_cell(ws, f"B{r_bank+3}", f"a/n {bank_an}", bold=True, size=13)

    wb.save(str(out_path))
    return str(out_path)
